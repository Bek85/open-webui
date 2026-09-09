"""Bounded, data-only rendering. No code execution, file inputs or remote resources."""

import math
import re
from io import BytesIO
from pathlib import Path
from typing import Literal
from xml.sax.saxutils import escape, quoteattr

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from markdown_it import MarkdownIt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, ConfigDict, Field, model_validator
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

MIME = {
    'pdf': 'application/pdf',
    'docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    'xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
}
MAX_OUTPUT = 10 * 1024 * 1024
MARKDOWN = MarkdownIt('commonmark', {'html': False, 'maxNesting': 20}).enable('table')


class Sheet(BaseModel):
    model_config = ConfigDict(extra='forbid')
    name: str = Field(min_length=1, max_length=31)
    columns: list[str] = Field(min_length=1, max_length=20)
    rows: list[list[str | int | float | bool | None]] = Field(default_factory=list, max_length=1000)

    @model_validator(mode='after')
    def validate_cells(self):
        if re.search(r'[\\/*?:\[\]]', self.name) or self.name.startswith("'") or self.name.endswith("'"):
            raise ValueError('Invalid sheet name')
        for row in [self.columns, *self.rows]:
            if len(row) != len(self.columns):
                raise ValueError('Every row must match the column count')
            for value in row:
                if isinstance(value, str) and (len(value) > 2000 or bad_controls(value)):
                    raise ValueError('Cell text is invalid or too long')
                if isinstance(value, float) and not math.isfinite(value):
                    raise ValueError('Nonfinite numbers are not supported')
        return self


class GenerateRequest(BaseModel):
    model_config = ConfigDict(extra='forbid')
    format: Literal['pdf', 'docx', 'xlsx']
    title: str = Field(min_length=1, max_length=160)
    content: str = Field(default='', max_length=60000)
    sheets: list[Sheet] = Field(default_factory=list, max_length=5)

    @model_validator(mode='after')
    def validate_document(self):
        if bad_controls(self.title + self.content) or not self.title.strip():
            raise ValueError('Invalid text')
        if self.format == 'xlsx':
            if not self.sheets or self.content:
                raise ValueError('Spreadsheets need sheets, not document content')
            if len({s.name.casefold() for s in self.sheets}) != len(self.sheets):
                raise ValueError('Sheet names must be unique')
            if sum((len(s.rows) + 1) * len(s.columns) for s in self.sheets) > 10000:
                raise ValueError('At most 10000 cells per workbook')
        elif not self.content.strip() or self.sheets:
            raise ValueError('Documents need content, not sheets')
        return self


def bad_controls(text):
    return bool(re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff]', text))


def filename_for(spec):
    stem = re.sub(r'[^\w -]', '_', spec.title, flags=re.UNICODE).strip(' ._')[:80] or 'document'
    return f'{stem}.{spec.format}'


# Explicit token-state transitions are kept together for auditability.
def blocks(content):  # noqa: C901
    """Flatten supported Markdown blocks; image references never trigger a fetch."""
    tokens = MARKDOWN.parse(content)
    if len(tokens) > 5000:
        raise ValueError('Document has too many formatting elements')
    result, row, table = [], None, None
    heading, lists, first_item_paragraph = 0, [], False
    for token in tokens:
        if token.type in ('bullet_list_open', 'ordered_list_open'):
            lists.append('ul' if token.type == 'bullet_list_open' else 'ol')
        elif token.type in ('bullet_list_close', 'ordered_list_close'):
            lists.pop()
        elif token.type == 'list_item_open':
            first_item_paragraph = True
        elif token.type == 'list_item_close':
            first_item_paragraph = False
        elif token.type == 'heading_open':
            heading = int(token.tag[1])
        elif token.type == 'heading_close':
            heading = 0
        elif token.type == 'table_open':
            table = []
        elif token.type == 'tr_open':
            row = []
        elif token.type == 'tr_close':
            table.append(row)
            row = None
        elif token.type == 'table_close':
            if len(table) > 200 or any(len(r) > 10 for r in table):
                raise ValueError('Document table exceeds 200 rows or 10 columns')
            result.append(('table', table))
            table = None
        elif token.type == 'inline':
            if row is not None:
                row.append(token.children or [])
            else:
                kind = f'h{heading}' if heading else lists[-1] if lists and first_item_paragraph else 'p'
                result.append((kind, token.children or []))
                first_item_paragraph = False
        elif token.type in ('fence', 'code_block'):
            result.append(('text', token.content))
    return result


def inline_pdf(tokens):
    output, link = [], False
    for token in tokens:
        if token.type in ('text', 'code_inline', 'image'):
            output.append(escape(token.content))
        elif token.type in ('softbreak', 'hardbreak'):
            output.append('<br/>')
        elif token.type in ('strong_open', 'strong_close', 'em_open', 'em_close'):
            output.append(
                {'strong_open': '<b>', 'strong_close': '</b>', 'em_open': '<i>', 'em_close': '</i>'}[token.type]
            )
        elif token.type == 'link_open':
            href = token.attrGet('href') or ''
            link = href.startswith(('https://', 'http://')) and len(href) <= 2048
            if link:
                output.append(f'<a href={quoteattr(href)} color="#16467a">')
        elif token.type == 'link_close' and link:
            output.append('</a>')
            link = False
    return ''.join(output)


def inline_docx(paragraph, tokens):  # noqa: C901 - explicit supported inline token types
    bold = italic = False
    link = None
    for token in tokens:
        if token.type == 'strong_open':
            bold = True
        elif token.type == 'strong_close':
            bold = False
        elif token.type == 'em_open':
            italic = True
        elif token.type == 'em_close':
            italic = False
        elif token.type == 'link_open':
            href = token.attrGet('href') or ''
            if href.startswith(('https://', 'http://')) and len(href) <= 2048:
                from docx.opc.constants import RELATIONSHIP_TYPE

                link = OxmlElement('w:hyperlink')
                link.set(qn('r:id'), paragraph.part.relate_to(href, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
                paragraph._p.append(link)
        elif token.type == 'link_close':
            link = None
        elif token.type in ('text', 'code_inline', 'image', 'softbreak', 'hardbreak'):
            run = paragraph.add_run('\n' if token.type.endswith('break') else token.content)
            run.bold, run.italic = bold, italic
            if link is not None:
                link.append(run._r)


def render_docx(spec):
    doc = Document()
    doc.styles['Normal'].font.name = 'Arial'
    doc.styles['Normal'].font.size = Pt(11)
    doc.core_properties.title = spec.title
    doc.core_properties.author = 'Prokuratura AI'
    doc.add_heading(spec.title, 0)
    for kind, data in blocks(spec.content):
        if kind == 'table':
            table = doc.add_table(rows=0, cols=len(data[0]))
            table.style = 'Light Shading Accent 1'
            for row in data:
                cells = table.add_row().cells
                for cell, tokens in zip(cells, row):
                    inline_docx(cell.paragraphs[0], tokens)
        elif kind == 'text':
            doc.add_paragraph(data)
        else:
            paragraph = (
                doc.add_heading(level=min(int(kind[1]), 4))
                if kind.startswith('h')
                else doc.add_paragraph(style={'ul': 'List Bullet', 'ol': 'List Number'}.get(kind))
            )
            inline_docx(paragraph, data)
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def register_fonts():
    base = Path('/usr/share/fonts/truetype/dejavu')
    for suffix, file in [
        ('', 'DejaVuSans.ttf'),
        ('-Bold', 'DejaVuSans-Bold.ttf'),
        ('-Oblique', 'DejaVuSans-Oblique.ttf'),
        ('-BoldOblique', 'DejaVuSans-BoldOblique.ttf'),
    ]:
        pdfmetrics.registerFont(TTFont('Document' + suffix, str(base / file)))
    pdfmetrics.registerFontFamily(
        'Document',
        normal='Document',
        bold='Document-Bold',
        italic='Document-Oblique',
        boldItalic='Document-BoldOblique',
    )


def render_pdf(spec):
    styles = {
        'p': ParagraphStyle('body', fontName='Document', fontSize=10, leading=15, spaceAfter=8, alignment=TA_LEFT),
        'heading': ParagraphStyle(
            'heading', fontName='Document-Bold', fontSize=14, leading=19, spaceBefore=10, spaceAfter=8
        ),
    }
    story = [Paragraph(escape(spec.title), styles['heading']), Spacer(1, 10)]
    number = 0
    for kind, data in blocks(spec.content):
        number = number + 1 if kind == 'ol' else 0
        if kind == 'table':
            cells = [[Paragraph(inline_pdf(cell), styles['p']) for cell in row] for row in data]
            table = LongTable(cells, colWidths=[(A4[0] - 90) / len(data[0])] * len(data[0]), repeatRows=1)
            table.setStyle(
                TableStyle(
                    [
                        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#e7eef7')),
                        ('GRID', (0, 0), (-1, -1), 0.3, colors.lightgrey),
                        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                    ]
                )
            )
            story.extend([table, Spacer(1, 10)])
        else:
            text = escape(data).replace('\n', '<br/>') if kind == 'text' else inline_pdf(data)
            if kind in ('ul', 'ol'):
                text = ('• ' if kind == 'ul' else f'{number}. ') + text
            story.append(Paragraph(text, styles['heading'] if kind.startswith('h') else styles['p']))
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=45,
        leftMargin=45,
        topMargin=45,
        bottomMargin=45,
        title=spec.title,
        author='Prokuratura AI',
    )
    doc.build(story)
    return buffer.getvalue()


def render_xlsx(spec):
    book = Workbook()
    book.remove(book.active)
    for sheet in spec.sheets:
        ws = book.create_sheet(sheet.name)
        for row_index, row in enumerate([sheet.columns, *sheet.rows], 1):
            for col_index, value in enumerate(row, 1):
                cell = ws.cell(row_index, col_index, value)
                # Even '=HYPERLINK(...)' stays literal text. No formulas, macros or external links.
                if isinstance(value, str):
                    cell.data_type = 's'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                if row_index == 1:
                    cell.font = Font(bold=True, color='FFFFFF')
                    cell.fill = PatternFill('solid', fgColor='16467A')
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for idx, title in enumerate(sheet.columns, 1):
            ws.column_dimensions[get_column_letter(idx)].width = min(45, max(18, len(title) + 4))
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def render(spec):
    result = {'docx': render_docx, 'pdf': render_pdf, 'xlsx': render_xlsx}[spec.format](spec)
    if len(result) > MAX_OUTPUT:
        raise ValueError('Generated file exceeds 10 MB')
    return result
