import unittest
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from openpyxl import load_workbook
from pydantic import ValidationError
from renderer import GenerateRequest, blocks, filename_for, register_fonts, render

TEXT = (
    '# Ҳуқуқий таҳлил\n\nЎзбекистон: Ғ, Қ, Ҳ, Ў. O‘zbekiston. Русский текст.\n\n'
    '**Қалин** ва *қийшиқ*. [Манба](https://lex.uz/docs/1)\n\n- Биринчи\n- Иккинчи\n\n'
    '| Номи | Сони |\n| --- | --- |\n| Мурожаат | 12 |'
)


class RendererTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        register_fonts()

    def test_docx_content_tables_links_and_lists(self):
        spec = GenerateRequest(format='docx', title='Синов', content=TEXT)
        output = render(spec)
        document = Document(BytesIO(output))
        self.assertIn('Ҳуқуқий таҳлил', '\n'.join(p.text for p in document.paragraphs))
        self.assertEqual(document.tables[0].cell(1, 1).text, '12')
        self.assertTrue(any(p.style.name == 'List Bullet' for p in document.paragraphs))
        with ZipFile(BytesIO(output)) as archive:
            self.assertIn(b'https://lex.uz/docs/1', archive.read('word/_rels/document.xml.rels'))

    def test_pdf_unicode_embedded_font(self):
        output = render(GenerateRequest(format='pdf', title='Синов', content=TEXT))
        self.assertTrue(output.startswith(b'%PDF-'))
        self.assertIn(b'/ToUnicode', output)
        self.assertIn(b'DejaVuSans', output)
        self.assertIn(b'https://lex.uz/docs/1', output)

    def test_xlsx_formula_like_cells_remain_text_and_numbers_remain_numbers(self):
        values = ['=HYPERLINK("http://bad.test","click")', '+SUM(1,2)', '@SUM(1,2)', '-1+2']
        spec = GenerateRequest(
            format='xlsx',
            title='Жадвал',
            sheets=[{'name': 'Ҳисобот', 'columns': ['Матн', 'Сон'], 'rows': [[v, 12] for v in values]}],
        )
        output = render(spec)
        book = load_workbook(BytesIO(output), data_only=False)
        sheet = book.active
        for idx, value in enumerate(values, 2):
            self.assertEqual(sheet.cell(idx, 1).value, value)
            self.assertEqual(sheet.cell(idx, 1).data_type, 's')
            self.assertEqual(sheet.cell(idx, 2).value, 12)
        self.assertEqual(sheet.freeze_panes, 'A2')
        self.assertFalse(book._external_links)

    def test_invalid_inputs_rejected(self):
        for payload in [
            {'format': 'html', 'title': 'x', 'content': 'x'},
            {'format': 'docx', 'title': 'x', 'content': 'x' * 60001},
            {'format': 'pdf', 'title': 'x', 'content': '\x00'},
            {'format': 'xlsx', 'title': 'x', 'sheets': [{'name': '../bad', 'columns': ['a'], 'rows': []}]},
            {'format': 'xlsx', 'title': 'x', 'sheets': [{'name': 'a', 'columns': ['a'], 'rows': [[1, 2]]}]},
            {'format': 'xlsx', 'title': 'x', 'sheets': [{'name': 'a', 'columns': ['a'], 'rows': [[float('nan')]]}]},
        ]:
            with self.subTest(payload=str(payload)[:80]), self.assertRaises(ValidationError):
                GenerateRequest.model_validate(payload)

    def test_no_path_or_image_loading(self):
        spec = GenerateRequest(
            format='pdf', title='../../secret', content='![alt](file:///etc/passwd)\n\n<img src="http://example.test">'
        )
        self.assertNotIn('/', filename_for(spec))
        self.assertTrue(render(spec).startswith(b'%PDF-'))
        self.assertEqual([kind for kind, _ in blocks('- one\n- two')], ['ul', 'ul'])


if __name__ == '__main__':
    unittest.main()
